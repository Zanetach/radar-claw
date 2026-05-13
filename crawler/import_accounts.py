from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .models import SourceAccount
from .parsing import row_to_accounts


def load_accounts_from_excel(path: Path) -> list[SourceAccount]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() for value in next(rows)]

    accounts: list[SourceAccount] = []
    for excel_row_number, row in enumerate(rows, start=2):
        accounts.extend(row_to_accounts(headers, row, excel_row_number))
    return accounts
